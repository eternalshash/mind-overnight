import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os
import re
import subprocess

# ----------------- STYLING CONSTANTS -----------------
COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_NAVY = RGBColor(0, 31, 63)       # #001F3F Dark Blue / Navy
COLOR_MUTED = RGBColor(70, 70, 70)
HEX_NAVY = "001F3F"
HEX_LIGHT_BG = "F4F6F9"
HEX_BORDER = "D0D7DE"
HEX_WHITE = "FFFFFF"

# ----------------- DOCUMENT SETUP -----------------
doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

style_normal = doc.styles['Normal']
style_normal.font.name = 'Times New Roman'
style_normal.font.size = Pt(11.5)
style_normal.font.color.rgb = COLOR_BLACK
style_normal.paragraph_format.line_spacing = 1.15
style_normal.paragraph_format.space_after = Pt(5)

# Helper functions
def add_header(title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(15.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_NAVY

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_before = Pt(1)
    p2.paragraph_format.space_after = Pt(8)
    r2 = p2.add_run(subtitle)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11.5)
    r2.font.italic = True
    r2.font.color.rgb = COLOR_MUTED

def add_meta(lines):
    for label, val in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(label + ": ")
        r1.font.name = 'Times New Roman'
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_NAVY
        r2 = p.add_run(val)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_BLACK

def add_sec_title(title_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title_str)
    r.font.name = 'Times New Roman'
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_NAVY

def add_sub_title(title_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title_str)
    r.font.name = 'Times New Roman'
    r.font.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = COLOR_NAVY

def add_para(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    r.font.color.rgb = COLOR_BLACK
    return p

def add_bullet(bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r1 = p.add_run(bold_prefix + " ")
        r1.font.name = 'Times New Roman'
        r1.font.bold = True
        r1.font.size = Pt(11.5)
        r1.font.color.rgb = COLOR_NAVY
    r2 = p.add_run(text)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11.5)
    r2.font.color.rgb = COLOR_BLACK
    return p

def add_code_block(code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)

    cell = table.cell(0, 0)
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_LIGHT_BG}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="{HEX_NAVY}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)

    mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="120" w:type="dxa"/>'
        f'<w:left w:w="180" w:type="dxa"/>'
        f'<w:bottom w:w="120" w:type="dxa"/>'
        f'<w:right w:w="180" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    cell._tc.get_or_add_tcPr().append(mar)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code_text)
    r.font.name = 'Courier New'
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_BLACK

    # Spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_image_centered(img_path, width_in, caption_text=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_in))

    if caption_text:
        pc = doc.add_paragraph()
        pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pc.paragraph_format.space_before = Pt(1)
        pc.paragraph_format.space_after = Pt(6)
        rc = pc.add_run(caption_text)
        rc.font.name = 'Times New Roman'
        rc.font.size = Pt(10)
        rc.font.italic = True
        rc.font.color.rgb = COLOR_NAVY

def style_table(table, col_widths, headers, rows):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = Inches(col_widths[i])
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(title)
        r.font.name = 'Times New Roman'
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_NAVY}"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.add_row().cells
        bg_color = HEX_LIGHT_BG if r_idx % 2 == 1 else HEX_WHITE
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].width = Inches(col_widths[c_idx])
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_BLACK
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
            row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)

    # Borders
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            b = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
                f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
                f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
                f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(b)

    # Spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

print("Framework initialized successfully.")
