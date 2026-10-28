import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os
import subprocess

doc = docx.Document()

# 1-inch margins all around
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Configure Default Normal Style
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(11.5)
font.color.rgb = RGBColor(0, 0, 0)

# Color Constants
COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_NAVY = RGBColor(0, 31, 63)      # #001F3F
HEX_NAVY = "001F3F"
HEX_BORDER_GRAY = "D0D7DE"
HEX_BG_LIGHT = "F8F9FA"
HEX_TH_BG = "001F3F"

def add_header(title, subtitle=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = COLOR_NAVY
    
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(8)
        r2 = p2.add_run(subtitle)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11)
        r2.font.italic = True
        r2.font.color.rgb = RGBColor(60, 60, 60)

def add_section_heading(title_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title_str)
    r.font.name = 'Times New Roman'
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_NAVY

def add_sub_heading(title_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title_str)
    r.font.name = 'Times New Roman'
    r.font.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = COLOR_NAVY

def add_p(text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.bold = True
        r_pre.font.size = Pt(11.5)
        r_pre.font.color.rgb = COLOR_BLACK
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    r.font.italic = italic
    r.font.color.rgb = COLOR_BLACK
    return p

def add_code_block(code_str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), HEX_BG_LIGHT))
    cell._tc.get_or_add_tcPr().append(shading_elm)
    
    borders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{}"/>
            <w:left w:val="single" w:sz="16" w:space="0" w:color="{}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{}"/>
        </w:tcBorders>
    '''.format(nsdecls('w'), HEX_BORDER_GRAY, HEX_NAVY, HEX_BORDER_GRAY, HEX_BORDER_GRAY))
    cell._tc.get_or_add_tcPr().append(borders)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(4)
    cp.paragraph_format.line_spacing = 1.05
    c_run = cp.add_run(code_str.strip())
    c_run.font.name = 'Menlo'
    c_run.font.size = Pt(8.5)
    c_run.font.color.rgb = RGBColor(20, 20, 20)
    
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)

def add_image_box(img_path, width_in=6.2):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width_in))

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def style_table(table, col_widths, headers, data_rows):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], HEX_TH_BG)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.runs[0]
        r.font.name = 'Times New Roman'
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
    
    # Data rows
    for r_idx, row_data in enumerate(data_rows):
        row_cells = table.add_row().cells
        bg_col = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if len(p.runs) > 0:
                p.runs[0].font.name = 'Times New Roman'
                p.runs[0].font.size = Pt(9.5)
                p.runs[0].font.color.rgb = COLOR_BLACK

    # Set column widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

print("Helper definitions successfully configured.")
