import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
import subprocess

doc = docx.Document()

# 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

style = doc.styles['Normal']
font = style.font
font.name = 'Helvetica'
font.size = Pt(10.5)
font.color.rgb = RGBColor(30, 30, 30)

def add_header(title, subtitle=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(20, 20, 20)
    
    if subtitle:
        p2 = doc.add_paragraph()
        r2 = p2.add_run(subtitle)
        r2.font.size = Pt(11)
        r2.font.italic = True
        r2.font.color.rgb = RGBColor(80, 80, 80)
    doc.add_paragraph()

def add_section_heading(title_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title_str)
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 51, 102)

def add_code_block(code_str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    shading_elm = parse_xml(r'<w:shd {} w:fill="F4F6F9"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading_elm)
    
    borders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>
            <w:left w:val="single" w:sz="16" w:space="0" w:color="0055AA"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>
        </w:tcBorders>
    '''.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(borders)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(4)
    c_run = cp.add_run(code_str.strip())
    c_run.font.name = 'Menlo'
    c_run.font.size = Pt(8.5)
    c_run.font.color.rgb = RGBColor(40, 40, 40)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_box(img_path):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(6.2))

# ==================== BUILD REPORT ====================

add_header("ECE 449 / 590: Object-Oriented Programming and Machine Learning\nProject 3 Report: Matrix and Tensor Operations for EasyNN in C++", 
           "Name: Shashwat Choudhry  |  CWID: A20507277  |  Semester: Fall 2026")

# ----------------- SECTION 1 -----------------
add_section_heading("1. Project Overview & Computational Model")

p1 = doc.add_paragraph()
p1.add_run(
    "In Project 3, we expand the C++ EasyNN runtime from purely scalar arithmetic to full multi-dimensional "
    "tensor and 2D matrix operations. The Python frontend builds computation graphs as Static Single Assignment "
    "(SSA) Directed Acyclic Graphs (DAGs) and passes both scalar and tensor operands down to libeasynn.so via ctypes.\n\n"
    "Key architectural objectives for this project include:\n"
    "• Unified Data Representation: Implementing a robust, lightweight C++ tensor structure that represents scalars, "
    "vectors (1D), matrices (2D), and higher-order tensors (3D, 4D) uniformly in contiguous row-major memory.\n"
    "• Tensor Input & Parameter Ingestion: Ingesting multi-dimensional arrays from Python NumPy buffers via add_op_param_ndarray "
    "and add_kwargs_ndarray without memory corruption or dangling pointers.\n"
    "• Element-Wise Arithmetic & Broadcasting: Supporting strict element-wise addition and subtraction on tensors of matching shapes, "
    "while rejecting invalid mixed-rank combinations.\n"
    "• Multi-Modal Multiplication: Dynamically branching between scalar multiplication, scalar-tensor scaling, and high-performance "
    "2D matrix multiplication (matmul).\n"
    "• Zero-Copy/Safe Memory Return: Returning multi-dimensional tensor results back across the C foreign function interface (FFI) "
    "using dim, shape, and data pointers that persist throughout the evaluation lifecycle."
)

# ----------------- SECTION 2 -----------------
add_section_heading("2. Data Structure Design: struct tensor")

p2 = doc.add_paragraph()
p2.add_run(
    "To handle both scalars and arbitrary n-dimensional tensors without polymorphism overhead, we implemented struct tensor in src/tensor.h:\n"
    "• shape: std::vector<size_t> storing the size of each axis. A scalar is cleanly identified by an empty shape vector (shape.empty() == true).\n"
    "• data: std::vector<double> holding the contiguous flattened elements in C-order (row-major).\n"
    "• Constructors:\n"
    "  - Default constructor: Creates an empty tensor.\n"
    "  - Scalar constructor: Takes a single double value, leaves shape empty, and stores data = {val}.\n"
    "  - Pre-allocated shape constructor: Allocates product(shape) elements initialized to 0.0.\n"
    "  - C-array constructor: Takes (dim, shape[], data[]). When dim == 0, it constructs a scalar. Otherwise, it copies dim dimensions and product(shape) contiguous values from the ctypes buffer into owned memory.\n"
    "• Helper methods: is_scalar() checks shape.empty(), dim() returns static_cast<int>(shape.size()), and size() returns data.size()."
)

add_code_block("""// src/tensor.h
struct tensor {
    std::vector<size_t> shape;
    std::vector<double> data;

    tensor() : shape{}, data{} {}
    tensor(double val) : shape{}, data{val} {}

    tensor(const std::vector<size_t> &s) : shape(s) {
        size_t total = 1;
        for (size_t d : s) total *= d;
        data.assign(total, 0.0);
    }

    tensor(int dim, const size_t s[], const double *d) {
        if (dim == 0) {
            shape.clear();
            if (d) data.assign(d, d + 1);
            else data.assign(1, 0.0);
        } else {
            shape.assign(s, s + dim);
            size_t total = 1;
            for (int i = 0; i < dim; ++i) total *= s[i];
            if (d) data.assign(d, d + total);
            else data.assign(total, 0.0);
        }
    }

    bool is_scalar() const { return shape.empty(); }
    int dim() const { return static_cast<int>(shape.size()); }
    size_t size() const { return data.size(); }
};""")

# ----------------- SECTION 3 -----------------
add_section_heading("3. Function Implementations & Evaluation Logic")

p3 = doc.add_paragraph()
p3.add_run(
    "The core functionality is distributed across the following functions in src/:\n\n"
    "1. add_op_param_ndarray (program.cpp & expression.cpp):\n"
    "   Forwards tensor parameters (such as Const tensors) from the program builder to the latest appended expression. "
    "   The expression instantiates a tensor object and stores it in op_param_tensor_[key].\n\n"
    "2. add_kwargs_ndarray (evaluation.cpp):\n"
    "   Called during evaluation setup when input variables are multi-dimensional arrays. Constructs a tensor from the provided "
    "   dim, shape, and raw double pointer and maps it by input name in kwargs_tensor_.\n\n"
    "3. Input & Const Operations (evaluation.cpp):\n"
    "   - Input: Looks up the tensor in kwargs_tensor_ matching expr.op_name_. Returns -1 if missing.\n"
    "   - Const: Retrieves the pre-stored parameter from expr.op_param_tensor_['value'].\n\n"
    "4. Element-Wise Addition (Add) and Subtraction (Sub):\n"
    "   - Scalar-Scalar: Evaluates simple algebraic sum or difference.\n"
    "   - Tensor-Tensor: Verifies that both operands have identical shape (a.shape == b.shape). Computes element-wise result "
    "     res.data[i] = a.data[i] +/- b.data[i] over all elements.\n"
    "   - Mixed Rank: Mismatched scalar-tensor additions/subtractions are rejected with an error, matching EasyNN specification.\n\n"
    "5. Multiplication (Mul):\n"
    "   - Scalar * Scalar: Standard algebraic product.\n"
    "   - Scalar * Tensor & Tensor * Scalar: Broadcasts the scalar by multiplying every element of the tensor, preserving tensor shape.\n"
    "   - Matrix Multiplication (Tensor * Tensor): Validates that both operands are strictly 2D matrices (a.dim() == 2 && b.dim() == 2) "
    "     and that inner dimensions match (a.shape[1] == b.shape[0]). Pre-allocates a result matrix of shape {M, N} and performs "
    "     three-nested-loop matrix multiplication using IKJ loop ordering for optimal CPU cache spatial locality.\n\n"
    "6. Negation (Neg):\n"
    "   Applies unary negation res.data[i] = -a.data[i] across all elements for both scalars and tensors.\n\n"
    "7. Execution & Result Pointer Interface (execute in libeasynn.cpp):\n"
    "   Upon successful evaluation, extracts metadata from the evaluator's result_tensor_:\n"
    "   *p_dim = eval->get_result_dim();\n"
    "   *p_shape = eval->get_result_shape();\n"
    "   *p_data = eval->get_result_data();\n"
    "   When returning a scalar, *p_dim is set to 0, *p_shape is nullptr, and *p_data points to data[0]. For tensors, *p_shape and "
    "   *p_data point directly into the contiguous buffers owned by result_tensor_, ensuring safe reads in Python ctypes."
)

add_code_block("""// Matrix Multiplication Engine (src/evaluation.cpp)
if (a.dim() != 2 || b.dim() != 2) return -1;
if (a.shape[1] != b.shape[0]) return -1;

size_t M = a.shape[0];
size_t K = a.shape[1];
size_t N = b.shape[1];
tensor res({M, N});

for (size_t i = 0; i < M; ++i) {
    for (size_t k = 0; k < K; ++k) {
        double a_ik = a.data[i * K + k];
        for (size_t j = 0; j < N; ++j) {
            res.data[i * N + j] += a_ik * b.data[k * N + j];
        }
    }
}
val = res;""")

# ----------------- SECTION 4 -----------------
add_section_heading("4. Verification & Testing Evidence")

p4 = doc.add_paragraph()
p4.add_run(
    "The implementation was thoroughly validated across unit, grading, and remote continuous integration suites:\n\n"
    "• Grading Test Suite (grade_p3.py):\n"
    "  All 10 test functions passed 100% on the first run:\n"
    "  - Q1: Input tensor support for 1D (9,) and 2D (9, 9) shapes [PASSED]\n"
    "  - Q2: Const tensor support for 1D and 2D arrays [PASSED]\n"
    "  - Q3: Element-wise Add across 1D, 2D, 3D, and 4D tensors [PASSED]\n"
    "  - Q4: Element-wise Sub across 1D, 2D, 3D, and 4D tensors [PASSED]\n"
    "  - Q5: 2D Matrix multiplication (11, 12) x (12, 13) [PASSED]\n"
    "  - Q6: Scalar-tensor and tensor-scalar multiplication [PASSED]\n"
    "  - Q7: Compound expression: x + y*z (matmul followed by addition) [PASSED]\n"
    "  - Q8: Compound expression: (x - y)*z (subtraction followed by matmul) [PASSED]\n"
    "  - Q9: Chained matrix multiplication: x*y*z across 3 matrices [PASSED]\n"
    "  - Q10: Complex DAG combining matmul, Const tensors, subtraction, and addition [PASSED]\n"
    "  Final Test Score: 10/10 functions passed (90/90 points).\n\n"
    "• Regression Suite:\n"
    "  Regression checks against grade_p1.py (10/10) and grade_p2.py (10/10) confirmed zero regressions in scalar support.\n\n"
    "• Central Dasan Git Server Automated Verification:\n"
    "  Pushed in 3 staggered commits to the central Dasan Git repository. The remote post-receive hook compiled and graded the submission, "
    "  confirming all functions passed remotely on Dasan."
)

add_image_box("/Users/schoudhry/Desktop/ece449_p3_tests_passed.png")
add_image_box("/Users/schoudhry/Desktop/ece449_p3_git_push_dasan.png")
add_image_box("/Users/schoudhry/Desktop/ece449_p3_easynn_test.png")

out_docx = "/Users/schoudhry/Desktop/ECE449_Project3_Report_ShashwatChoudhry.docx"
doc.save(out_docx)
print(f"Generated DOCX: {out_docx}")

# Convert to Pages using AppleScript
out_pages = "/Users/schoudhry/Desktop/ECE449_Project3_Report_ShashwatChoudhry.pages"
applescript = f'''
tell application "Pages"
    activate
    open POSIX file "{out_docx}"
    delay 2
    set frontDoc to front document
    save frontDoc in POSIX file "{out_pages}"
end tell
'''
try:
    res = subprocess.run(["osascript", "-e", applescript], capture_output=True, text=True, timeout=30)
    print("Pages conversion output:", res.stdout, res.stderr)
    print(f"Generated Pages: {out_pages}")
except Exception as e:
    print(f"Pages conversion error: {e}")
