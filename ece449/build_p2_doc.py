import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = docx.Document()

# Page Margins (1 inch all around)
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

style = doc.styles['Normal']
font = style.font
font.name = 'Helvetica'
font.size = Pt(10.5)
font.color.rgb = RGBColor(30, 30, 30)

def add_header(title, subtitle=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(20, 20, 20)
    
    if subtitle:
        p2 = doc.add_paragraph()
        r2 = p2.add_run(subtitle)
        r2.font.size = Pt(11)
        r2.font.italic = True
        r2.font.color.rgb = RGBColor(80, 80, 80)
    doc.add_paragraph()

def add_section_heading(title_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title_str)
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 51, 102)

def add_code_block(code_str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    shading_elm = parse_xml(r'<w:shd {} w:fill="F4F6F9"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading_elm)
    
    borders = parse_xml(r'''
        <w:tcBorders {} >
            <w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>
            <w:left w:val="single" w:sz="16" w:space="0" w:color="0055AA"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>
        </w:tcBorders>
    '''.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(borders)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(4)
    c_run = cp.add_run(code_str.strip())
    c_run.font.name = 'Menlo'
    c_run.font.size = Pt(8.5)
    c_run.font.color.rgb = RGBColor(40, 40, 40)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_box(img_path, caption_str=""):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(6.2))
        
        if caption_str:
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(8)
            cpr = cp.add_run(caption_str)
            cpr.font.size = Pt(9)
            cpr.font.italic = True
            cpr.font.color.rgb = RGBColor(110, 110, 110)

# ==================== BUILD REPORT ====================

add_header("ECE 449 / 590: Object-Oriented Programming and Machine Learning\nProject 2 Report: Scalar Operations for EasyNN in C++", 
           "Name: Shashwat Choudhry  |  CWID: A20507277  |  Semester: Fall 2026")

# ----------------- SECTION 1 -----------------
add_section_heading("1. Project Overview & System Architecture")

p1 = doc.add_paragraph()
p1.add_run("In Project 2, we implement scalar operations for the EasyNN framework in C++, interfacing with Python through the ctypes foreign function library. Python constructs the neural network computation as a Directed Acyclic Graph (DAG) in Static Single Assignment (SSA) form, and compiles it by dispatching C foreign calls to libeasynn.so.\n\n"
           "A core design principle of the system is the strict architectural separation between program and evaluation:\n"
           "• program: Represents the static computation graph. It stores the immutable sequence of SSA expressions in topological evaluation order, along with static operator parameters (such as Const scalar values). Once constructed, it remains unmodified.\n"
           "• evaluation: Represents a specific dynamic execution instance of the graph. It stores runtime input arguments (kwargs) passed per inference run and computes intermediate results. This design allows a single program DAG to be evaluated repeatedly across multiple sets of inputs without mutating the underlying graph.")

# ----------------- SECTION 2 -----------------
add_section_heading("2. Function Implementations & Class Design")

p2 = doc.add_paragraph()
p2.add_run("The following C++ interface and member functions were implemented in src/:\n\n"
           "• program* create_program():\n"
           "  Allocates and returns a new heap-allocated program instance to serve as the DAG container for SSA expressions.\n\n"
           "• void append_expression(program *prog, int expr_id, const char *op_name, const char *op_type, int inputs[], int num_inputs):\n"
           "  Instantiates an expression object with its unique ID, operator name, operator type, and array of input operand IDs, and appends it to the program's internal expressions vector following topological order.\n\n"
           "• int add_op_param_double(program *prog, const char *key, double value):\n"
           "  Attaches a scalar parameter to the most recently appended expression. In Project 2, this is utilized by the Const operator to store key 'value' with its assigned double constant. Returns 0 for success.\n\n"
           "• int add_op_param_ndarray(program *prog, const char *key, int dim, size_t shape[], double data[]):\n"
           "  Stubbed for future tensor support in Project 4; returns 0 for success.\n\n"
           "• evaluation* build(program *prog):\n"
           "  Constructs an evaluation instance, transferring the program's expression vector into the evaluator, and returns the evaluation pointer back to Python.\n\n"
           "• void add_kwargs_double(evaluation *eval, const char *key, double value):\n"
           "  Stores scalar input variables into the evaluation object's kwargs_double_ map, associating input variable names (e.g., 'x', 'y', 'z') with their runtime numerical values.\n\n"
           "• void add_kwargs_ndarray(evaluation *eval, const char *key, int dim, size_t shape[], double data[]):\n"
           "  Stubbed for future tensor inputs in Project 4; returns cleanly.\n\n"
           "• int execute(evaluation *eval, int *p_dim, size_t **p_shape, double **p_data):\n"
           "  Executes the evaluation pipeline, sets *p_dim = 0, *p_shape = nullptr, and assigns *p_data = &eval->get_result() to pass the final scalar result back to Python ctypes.\n\n"
           "• int evaluation::execute():\n"
           "  The core computational engine. Iterates sequentially through the topologically sorted expressions using an internal map<int, double> val_map to resolve operand values:\n"
           "  - Input: Looks up the operand value by op_name in kwargs_double_.\n"
           "  - Const: Looks up the literal constant from op_param_double_['value'].\n"
           "  - Add: Computes val_map[inputs[0]] + val_map[inputs[1]].\n"
           "  - Sub: Computes val_map[inputs[0]] - val_map[inputs[1]].\n"
           "  - Mul: Computes val_map[inputs[0]] * val_map[inputs[1]].\n"
           "  - Neg: Computes -val_map[inputs[0]].\n"
           "  Stores each computed value in val_map[expr_id] and updates result_ with the last evaluated expression.")

add_code_block("""// Core Evaluation Engine (src/evaluation.cpp)
int evaluation::execute() {
    std::map<int, double> val_map;
    for (const auto &expr : expressions_) {
        double val = 0.0;
        const std::string &op_type = expr.op_type_;

        if (op_type == "Input") {
            auto it = kwargs_double_.find(expr.op_name_);
            if (it == kwargs_double_.end()) return -1;
            val = it->second;
        } else if (op_type == "Const") {
            auto it = expr.op_param_double_.find("value");
            if (it == expr.op_param_double_.end()) return -1;
            val = it->second;
        } else if (op_type == "Add") {
            assert(expr.inputs_.size() >= 2);
            val = val_map[expr.inputs_[0]] + val_map[expr.inputs_[1]];
        } else if (op_type == "Sub") {
            assert(expr.inputs_.size() >= 2);
            val = val_map[expr.inputs_[0]] - val_map[expr.inputs_[1]];
        } else if (op_type == "Mul") {
            assert(expr.inputs_.size() >= 2);
            val = val_map[expr.inputs_[0]] * val_map[expr.inputs_[1]];
        } else if (op_type == "Neg") {
            assert(expr.inputs_.size() >= 1);
            val = -val_map[expr.inputs_[0]];
        } else {
            return -1;
        }
        val_map[expr.expr_id_] = val;
        result_ = val;
    }
    return 0;
}""")

# ----------------- SECTION 3 -----------------
add_section_heading("3. Verification & Grading Suite Results")

p3 = doc.add_paragraph()
p3.add_run("The implementation was verified against the provided test drivers, the local grading suite (grade_p2.py), and the central Dasan Git server automated grading pipeline:\n\n"
           "1. Standalone C++ Driver (easynn_test):\n"
           "   Constructs a DAG with an Input 'a' and an Add operator summing 'a + a'. With input a = 5.0, the C++ execution cleanly produced res = 10.000000.\n\n"
           "2. Grading Test Suite (grade_p2.py):\n"
           "   All 10 grading questions passed with 100% accuracy, verifying scalar inputs, constants, multi-operand expressions, operator precedence, and repeated evaluation with randomized kwargs across 10 iterations:\n"
           "   • Q1: Input('x') [PASSED]\n"
           "   • Q2: Const(c) [PASSED]\n"
           "   • Q3: x + c [PASSED]\n"
           "   • Q4: x * c [PASSED]\n"
           "   • Q5: x + y [PASSED]\n"
           "   • Q6: x * y [PASSED]\n"
           "   • Q7: x + y * z [PASSED]\n"
           "   • Q8: (x - y) * z [PASSED]\n"
           "   • Q9: x*y - x*c - y*c + c*c (10 random iterations) [PASSED]\n"
           "   • Q10: x*y - x*c - y*c + c*c (10 random iterations) [PASSED]\n"
           "   Final Result: 10 functions passed (90/90 test points).\n\n"
           "3. Central Dasan Git Server Automated Grading:\n"
           "   Changes were committed and pushed across 3 staggered commits. The central Dasan Git repository post-receive hook executed automated testing, confirming all 10 functions passed remotely on Dasan.")

add_image_box("/Users/schoudhry/Desktop/ece449_p2_git_push_dasan.png", 
              "Figure 3.1: Central Dasan Git repository push results showing remote execution passing all 10 functions.")

add_image_box("/Users/schoudhry/Desktop/ece449_p2_tests_passed.png", 
              "Figure 3.2: Complete VS Code terminal output showing successful build and all 10 grading test cases passed.")

add_image_box("/Users/schoudhry/Desktop/ece449_p2_easynn_test.png", 
              "Figure 3.3: Standalone C++ testing program (easynn_test) execution verifying DAG evaluation.")

out_docx = "/Users/schoudhry/Desktop/ECE449_Project2_Report_ShashwatChoudhry.docx"
doc.save(out_docx)
print(f"Project 2 Report successfully built: {out_docx}")
